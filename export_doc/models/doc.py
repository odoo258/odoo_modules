# -*- coding: utf-8 -*-

import functools
import logging
import werkzeug
import json
from docx import Document
import re
import datetime
from cStringIO import StringIO

from odoo.tools.misc import str2bool, xlwt
from odoo.addons.web.controllers.main import Export as Export
from odoo.addons.web.controllers.main import ExportFormat as ExportFormat
from odoo import http
from odoo.http import content_disposition, dispatch_rpc, request, \
                      serialize_exception as _serialize_exception

_logger = logging.getLogger(__name__)


def serialize_exception(f):
    @functools.wraps(f)
    def wrap(*args, **kwargs):
        try:
            return f(*args, **kwargs)
        except Exception, e:
            _logger.exception("An exception occured during an http request")
            se = _serialize_exception(e)
            error = {
                'code': 200,
                'message': "Odoo Server Error",
                'data': se
            }
            return werkzeug.exceptions.InternalServerError(json.dumps(error))
    return wrap


class ExportDocFormat(Export):

    @http.route('/web/export/formats', type='json', auth="user")
    def formats(self):
        res = super(ExportDocFormat, self).formats()
        res.append({'tag': 'doc', 'label': 'DOC'})
        return res


class DocExport(ExportFormat, http.Controller):
    raw_data = True

    @http.route('/web/export/doc', type='http', auth="user")
    @serialize_exception
    def index(self, data, token):
        return self.base(data, token)

    @property
    def content_type(self):
        return 'application/vnd.ms-word'

    def filename(self, base):
        return base + '.doc'

    def from_data(self, fields, rows):
        document = Document()
        table = document.add_table(rows=len(rows)+1, cols=len(fields)+1)
        heading_cells = table.rows[0].cells
        heading_cells[0].text = "N"
        for i, fieldname in enumerate(fields):
            heading_cells[i+1].text = fieldname

        index = 1
        for row_index, row in enumerate(rows):
            cell = table.cell(row_index+1, 0)
            cell.text = str(index)
            index += 1
            for cell_index, cell_value in enumerate(row):
                cell = table.cell(row_index+1, cell_index+1)
                cell.text = str(cell_value)

        fp = StringIO()
        document.save(fp)
        fp.seek(0)
        data = fp.read()
        fp.close()
        return data
